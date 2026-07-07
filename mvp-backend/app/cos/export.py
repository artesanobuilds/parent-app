"""Phase 3b — regenerate the family plan as a workbook-style .xlsx.

Renders the chief-of-staff plan (tracked activities, calendar, logistics/coverage,
to-dos) plus the derived intelligence from `cos/planning.build_plan_summary`
(cost tally, conflicts, coverage gaps, uncovered weeks) into a multi-sheet
openpyxl workbook — the machine-generated cousin of the hand-built Alvarado
workbook. Preserves the pending/confirmed discipline via cell coloring, and
flags coverage gaps in red.

Pure function: takes already-fetched rows, returns xlsx bytes. The router does
the storage reads + streaming so this stays trivially testable.
"""
from __future__ import annotations

import io
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Hearth-ish palette (ARGB).
_HEADER_FILL = PatternFill("solid", fgColor="4F5D3A")   # olive
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_CONFIRMED_FILL = PatternFill("solid", fgColor="E8EBDC")  # olive tint
_PENDING_FILL = PatternFill("solid", fgColor="FBEED0")    # warm/amber
_GAP_FILL = PatternFill("solid", fgColor="F6E4E4")        # red tint (coverage gap)
_GAP_FONT = Font(color="A03A3A", bold=True)
_TITLE_FONT = Font(size=14, bold=True, color="1A1A1A")


def _sheet(wb: Workbook, title: str, headers: List[str]):
    ws = wb.create_sheet(title[:31])  # Excel caps sheet names at 31 chars
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    return ws


def _autosize(ws) -> None:
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width + 2, 10), 48)


def _confirmedish(status: Optional[str], price_status: Optional[str] = None) -> Optional[bool]:
    s = (status or "").lower()
    if s in ("confirmed", "signed_up", "done"):
        return True
    if "pending" in s or s in ("exploring", "interested", "tbd", "waitlisted") or price_status == "tbd":
        return False
    return None


def build_plan_workbook(
    *,
    tracked: List[Dict[str, Any]],
    calendar: List[Dict[str, Any]],
    recurring: List[Dict[str, Any]],
    caregivers: List[Dict[str, Any]],
    logistics: List[Dict[str, Any]],
    todos: List[Dict[str, Any]],
    summary: Dict[str, Any],
    family_name: str = "Family",
) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)  # drop the default sheet; we name our own

    # ── Overview (cost tally + counts) ───────────────────────────────
    ov = wb.create_sheet("Overview")
    ov["A1"] = f"{family_name} — Activity Plan"; ov["A1"].font = _TITLE_FONT
    cost = (summary or {}).get("cost") or {}
    rows = [
        ("Confirmed total", f"${cost.get('confirmed_total', 0):,.0f}"),
        ("Items TBD (price)", cost.get("tbd_count", 0)),
        ("Tracked activities", cost.get("item_count", len(tracked))),
        ("Coverage gaps", len((summary or {}).get("coverage_gaps") or [])),
        ("Time conflicts", len((summary or {}).get("conflicts") or [])),
        ("Uncovered weeks", len((summary or {}).get("uncovered_weeks") or [])),
    ]
    for i, (k, v) in enumerate(rows, start=3):
        ov[f"A{i}"] = k; ov[f"B{i}"] = v; ov[f"A{i}"].font = Font(bold=True)
    by_kid = cost.get("by_kid") or {}
    if by_kid:
        ov["A11"] = "Confirmed by kid"; ov["A11"].font = Font(bold=True)
        for j, (kid, amt) in enumerate(by_kid.items(), start=12):
            ov[f"A{j}"] = kid; ov[f"B{j}"] = f"${amt:,.0f}"
    _autosize(ov)

    # ── Activities (pending/confirmed coloring) ──────────────────────
    ws = _sheet(wb, "Activities", ["Kid", "Activity", "Provider", "Status", "Price", "Price status", "Week", "Signup"])
    for a in tracked:
        ws.append([
            a.get("kid_name"), a.get("name"), a.get("provider"), a.get("status"),
            (f"${a['price_amount']:,.0f}" if a.get("price_amount") is not None else None),
            a.get("price_status"), a.get("week_start"), a.get("signup_url"),
        ])
        conf = _confirmedish(a.get("status"), a.get("price_status"))
        if conf is not None:
            fill = _CONFIRMED_FILL if conf else _PENDING_FILL
            ws.cell(row=ws.max_row, column=4).fill = fill
    _autosize(ws)

    # ── Calendar (week × kid) ────────────────────────────────────────
    ws = _sheet(wb, "Calendar", ["Week", "Kid", "Title", "Kind", "Start", "End", "Status", "Location"])
    for e in sorted(calendar, key=lambda x: (x.get("week_start") or "", x.get("kid_name") or "")):
        ws.append([
            e.get("week_start"), e.get("kid_name"), e.get("title"), e.get("kind"),
            e.get("start_ts"), e.get("end_ts"), e.get("status"), e.get("location_text"),
        ])
        conf = _confirmedish(e.get("status"))
        if conf is not None:
            ws.cell(row=ws.max_row, column=7).fill = _CONFIRMED_FILL if conf else _PENDING_FILL
    _autosize(ws)

    # ── Logistics & coverage (gaps in red) ───────────────────────────
    cg_names = {c.get("id"): c.get("name") for c in caregivers}
    ws = _sheet(wb, "Logistics", ["Week", "Kid", "Task", "Caregiver", "Status", "Notes"])
    for lo in sorted(logistics, key=lambda x: (x.get("week_start") or "", x.get("kid_name") or "")):
        is_gap = (lo.get("status") == "gap") or (lo.get("status") != "assigned" and not lo.get("caregiver_id"))
        ws.append([
            lo.get("week_start"), lo.get("kid_name"), lo.get("task"),
            cg_names.get(lo.get("caregiver_id")) or ("— GAP —" if is_gap else None),
            lo.get("status"), lo.get("notes"),
        ])
        if is_gap:
            for c in range(1, 7):
                ws.cell(row=ws.max_row, column=c).fill = _GAP_FILL
            ws.cell(row=ws.max_row, column=5).font = _GAP_FONT
    _autosize(ws)

    # ── Caregivers + recurring (references) ──────────────────────────
    ws = _sheet(wb, "Caregivers", ["Name", "Role", "Notes"])
    for c in caregivers:
        ws.append([c.get("name"), c.get("role"), c.get("notes")])
    _autosize(ws)
    if recurring:
        ws = _sheet(wb, "Recurring", ["Kid", "Title", "Provider", "Weekday", "Start", "End", "Transporter"])
        _WD = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
        for r in recurring:
            wd = r.get("weekday")
            ws.append([
                r.get("kid_name"), r.get("title"), r.get("provider_name"),
                (_WD[wd] if isinstance(wd, int) and 0 <= wd < 7 else wd),
                r.get("start_time"), r.get("end_time"), r.get("transporter"),
            ])
        _autosize(ws)

    # ── To-dos ───────────────────────────────────────────────────────
    ws = _sheet(wb, "To-dos", ["Kid", "Title", "Detail", "Due", "Status"])
    for t in todos:
        ws.append([t.get("kid_name"), t.get("title"), t.get("detail"), t.get("due_date"), t.get("status")])
    _autosize(ws)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
