"""Phase 3b — plan workbook export (app/cos/export.py). No network; the builder
is a pure function over already-fetched rows, so we assert structure + that the
bytes are a valid, re-openable xlsx."""
from __future__ import annotations

import io
import sys
from pathlib import Path

BACKEND_ROOT = Path(__file__).resolve().parents[1]
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from openpyxl import load_workbook  # noqa: E402

from app.cos.export import build_plan_workbook  # noqa: E402


def _wb():
    xlsx = build_plan_workbook(
        tracked=[
            {"kid_name": "Ceiba", "name": "Art Camp", "status": "confirmed",
             "price_amount": 450, "price_status": "confirmed", "week_start": "2026-06-15"},
            {"kid_name": "Alkina", "name": "Swim", "status": "pending_eligibility",
             "price_amount": 200, "price_status": "tbd", "week_start": "2026-06-15"},
        ],
        calendar=[{"week_start": "2026-06-15", "kid_name": "Ceiba", "title": "Art Camp",
                   "kind": "camp", "status": "confirmed"}],
        recurring=[],
        caregivers=[{"id": "cg1", "name": "Nanny Sam", "role": "nanny"}],
        logistics=[{"week_start": "2026-06-29", "kid_name": "Ceiba", "task": "pickup",
                    "caregiver_id": None, "status": "gap"}],
        todos=[{"kid_name": "Alkina", "title": "Confirm eligibility", "status": "open"}],
        summary={"cost": {"confirmed_total": 450, "tbd_count": 1, "item_count": 2,
                          "by_kid": {"Ceiba": 450}},
                 "coverage_gaps": [{"reason": "x"}], "conflicts": [], "uncovered_weeks": []},
        family_name="Alvarado",
    )
    return xlsx, load_workbook(io.BytesIO(xlsx))


def test_workbook_is_valid_xlsx_with_expected_sheets():
    xlsx, wb = _wb()
    assert xlsx[:2] == b"PK"  # xlsx is a zip
    assert {"Overview", "Activities", "Calendar", "Logistics", "Caregivers", "To-dos"} <= set(wb.sheetnames)


def test_activities_and_cost_rendered():
    _, wb = _wb()
    ws = wb["Activities"]
    names = [ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)]
    assert "Art Camp" in names and "Swim" in names
    assert wb["Overview"]["A1"].value == "Alvarado — Activity Plan"


def test_coverage_gap_row_is_flagged_red():
    _, wb = _wb()
    ws = wb["Logistics"]
    # The single gap row's status cell should carry the red gap fill (non-default).
    fill = ws.cell(row=2, column=5).fill
    assert fill is not None and fill.fgColor is not None
    assert str(fill.fgColor.rgb).endswith("F6E4E4")
